from io import BytesIO
from datetime import datetime

import pytest
from openpyxl import load_workbook


@pytest.fixture
def deliveries(app, monkeypatch):
    captured = []
    service = app.extensions["wecom_service"]
    service.enabled = True

    def send(person, stream, filename, count):
        workbook = load_workbook(BytesIO(stream.getvalue()), data_only=True)
        captured.append([row[0] for row in workbook.active.iter_rows(min_row=5, values_only=True)])
        workbook.close()
        return {"person": person, "taskCount": count, "mentioned": False}

    monkeypatch.setattr(service, "send_excel", send)
    return captured


def test_stage_options_and_filter_apply_to_list_and_export(client):
    options = client.get("/api/config").get_json()["data"]["stages"]
    assert {"code": "APPEALING", "name": "申诉中"} in options
    assert not any(option["code"] == "CLOSED" for option in options)
    data = client.get("/api/tasks?scope=all&stage=APPEALING").get_json()["data"]
    assert {item["id"] for item in data["items"]} == {103}
    response = client.get("/api/export?scope=all&stage=APPEALING")
    workbook = load_workbook(BytesIO(response.data))
    assert workbook.active.max_row == 5
    assert workbook.active["A5"].value == "ADM-260909-003"
    workbook.close()
    assert client.get("/api/tasks?scope=all&stage=INVALID").status_code == 400


def test_wecom_sends_all_filtered_pages_not_all_person_tasks(client, deliveries):
    response = client.post("/api/wecom/send?scope=team&person=黄娜娟&source=CTRIP&pageSize=1",
                           json={"person": "黄娜娟", "mode": "filtered"})
    assert response.status_code == 200
    assert set(deliveries[0]) == {"ADM-260909-001", "ADM-260909-005"}


def test_wecom_sends_only_selected_ids(client, deliveries):
    response = client.post("/api/wecom/send?scope=team&person=黄娜娟",
                           json={"person": "黄娜娟", "mode": "selected", "admIds": [101, 104, 101]})
    assert response.status_code == 200
    assert set(deliveries[0]) == {"ADM-260909-001", "ADM-260909-004"}


@pytest.mark.parametrize("ids", [[], [9999], [107], [True], ["101"]])
def test_bad_or_out_of_scope_selection_stops_without_sending(client, deliveries, ids):
    response = client.post("/api/wecom/send?scope=team&person=黄娜娟",
                           json={"person": "黄娜娟", "mode": "selected", "admIds": ids})
    assert response.status_code == 400
    assert deliveries == []


def test_closed_selection_stops_without_sending(client, repository, deliveries):
    repository.rows[0]["adm_status"] = 3
    response = client.post("/api/wecom/send?scope=all",
                           json={"person": "黄娜娟", "mode": "selected", "admIds": [101]})
    assert response.status_code == 400
    assert deliveries == []


def test_empty_filter_and_missing_mode_do_not_send(client, deliveries):
    assert client.post("/api/wecom/send?scope=all&search=NOTFOUND",
                       json={"person": "黄娜娟", "mode": "filtered"}).status_code == 400
    assert client.post("/api/wecom/send?scope=all", json={"person": "黄娜娟"}).status_code == 400
    assert deliveries == []


def test_send_respects_stage_and_creation_date(client, repository, deliveries):
    repository.rows[2]["create_time"] = datetime(2026, 9, 10, 23, 59, 59)
    response = client.post("/api/wecom/send?scope=all&stage=APPEALING&createdFrom=2026-09-10&createdTo=2026-09-10",
                           json={"person": "黄娜娟", "mode": "filtered"})
    assert response.status_code == 200
    assert deliveries == [["ADM-260909-003"]]


def test_selection_outside_source_filter_is_not_sent(client, deliveries):
    response = client.post("/api/wecom/send?scope=all&source=CTRIP",
                           json={"person": "黄娜娟", "mode": "selected", "admIds": [101, 104]})
    assert response.status_code == 400
    assert deliveries == []


def test_send_limit_stops_before_delivery(client, repository, deliveries, monkeypatch):
    monkeypatch.setattr(repository, "list_tasks", lambda filters: {"items": [], "pagination": {"total": 10001}})
    response = client.post("/api/wecom/send?scope=all",
                           json={"person": "黄娜娟", "mode": "filtered"})
    assert response.status_code == 400
    assert deliveries == []
