from io import BytesIO

from openpyxl import load_workbook

from adm_app.adm_import_converter import IMPORT_COLUMNS


def _business_workbook(client, query="scope=team&person=黄娜娟&source=CTRIP"):
    response = client.get(f"/api/export?{query}")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    return workbook


def _set_business_values(workbook, adm_no: str, values: dict[str, object]):
    sheet = workbook["ADM待处理"]
    headers = {cell.value: cell.column for cell in sheet[4]}
    for row_number in range(5, sheet.max_row + 1):
        if str(sheet.cell(row_number, headers["ADM单号"]).value) == adm_no:
            for label, value in values.items():
                sheet.cell(row_number, headers[label]).value = value
            return
    raise AssertionError(f"未找到{adm_no}")


def _post_workbook(client, workbook):
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return client.post(
        "/api/adm-import/convert",
        data={"file": (stream, "业务回传.xlsx")},
        content_type="multipart/form-data",
    )


def _converted_row(response, adm_no: str):
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), data_only=True)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row_number in range(2, sheet.max_row + 1):
        if str(sheet.cell(row_number, headers["ADM单号（业务唯一标识）"]).value) == adm_no:
            return workbook, sheet, row_number, headers
    raise AssertionError(f"转换结果未找到{adm_no}")


def test_business_return_converts_to_complete_adm_import(client):
    workbook = _business_workbook(client)
    _set_business_values(workbook, "ADM-260909-001", {
        "差异说明": "业务确认：返点差异待处理",
        "是否确认": "资料不全",
        "实际责任人": "李志君",
        "处理进度": "已录入差异",
        "申诉状态": "不提交可结案",
        "申诉原因": "业务确认不申诉",
        "结案处理结果": "待ADM专员审核",
        "恢复编码": "PNR8X2",
    })

    response = _post_workbook(client, workbook)
    output, sheet, row_number, headers = _converted_row(response, "ADM-260909-001")

    assert [cell.value for cell in sheet[1]] == [label for _, label in IMPORT_COLUMNS]
    assert len(IMPORT_COLUMNS) == 38
    assert sheet.cell(row_number, headers["主键ID"]).value == 101
    assert sheet.cell(row_number, headers["实际责任人"]).value == "李志君"
    assert sheet.cell(row_number, headers["处理状态"]).value == "待审核"
    assert sheet.cell(row_number, headers["申诉状态"]).value == "不申诉"
    assert sheet.cell(row_number, headers["申诉结果"]).value is None
    assert sheet.cell(row_number, headers["细分差异原因（可备注，自由文本描述具体差异情况）"]).value == "业务确认：返点差异待处理"
    remark = sheet.cell(row_number, headers["备注"]).value
    assert "原系统备注" in remark
    assert "确认=资料不全" in remark
    assert "进度=已录入差异" in remark
    assert "恢复编码=PNR8X2" in remark
    assert sheet.auto_filter.ref.startswith("A1:AL")
    assert sheet.freeze_panes == "A2"
    output.close()


def test_appeal_result_becomes_pending_audit(client):
    workbook = _business_workbook(client)
    _set_business_values(workbook, "ADM-260909-001", {
        "申诉状态": "已提交",
        "申诉结果": "申诉成功",
    })
    response = _post_workbook(client, workbook)
    output, sheet, row_number, headers = _converted_row(response, "ADM-260909-001")
    assert sheet.cell(row_number, headers["处理状态"]).value == "待审核"
    assert sheet.cell(row_number, headers["申诉状态"]).value == "申诉"
    assert sheet.cell(row_number, headers["申诉结果"]).value == "申诉成功"
    output.close()


def test_missing_adm_stops_conversion(client):
    workbook = _business_workbook(client)
    _set_business_values(workbook, "ADM-260909-001", {"ADM单号": "ADM-NOT-FOUND"})
    response = _post_workbook(client, workbook)
    assert response.status_code == 400
    assert "数据库中未找到有效ADM" in response.get_json()["message"]

