from io import BytesIO

from openpyxl import load_workbook


def recovery_workbook(client, code="PNR8X2") -> BytesIO:
    response = client.get("/api/export?scope=team&person=黄娜娟&source=CTRIP")
    workbook = load_workbook(BytesIO(response.data))
    sheet = workbook["ADM待处理"]
    headers = {cell.value: cell.column for cell in sheet[4]}
    sheet.cell(5, headers["恢复编码"], code)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def test_import_and_follow_up_recovery_code(client):
    response = client.post(
        "/api/recovery/import",
        data={"operator": "曾芸芸", "file": (recovery_workbook(client), "ADM.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    result = response.get_json()["data"]
    assert result["importedCount"] == 1
    assert result["createdCount"] == 1

    listing = client.get("/api/recovery?status=PENDING").get_json()["data"]
    assert listing["pagination"]["total"] == 1
    item = listing["items"][0]
    assert item["recovery_code"] == "PNR8X2"
    assert item["statusName"] == "待恢复"

    updated = client.patch(
        f"/api/recovery/{item['id']}",
        json={"status": "COMPLETED", "remark": "编码已恢复", "handler": "曾芸芸"},
    )
    assert updated.status_code == 200
    assert updated.get_json()["data"]["statusName"] == "已完成"


def test_changed_code_resets_completed_record(client):
    first = recovery_workbook(client, "CODE01")
    client.post(
        "/api/recovery/import",
        data={"operator": "曾芸芸", "file": (first, "first.xlsx")},
        content_type="multipart/form-data",
    )
    item = client.get("/api/recovery").get_json()["data"]["items"][0]
    client.patch(
        f"/api/recovery/{item['id']}",
        json={"status": "COMPLETED", "remark": "完成", "handler": "曾芸芸"},
    )

    second = recovery_workbook(client, "CODE02")
    result = client.post(
        "/api/recovery/import",
        data={"operator": "曾芸芸", "file": (second, "second.xlsx")},
        content_type="multipart/form-data",
    ).get_json()["data"]
    assert result["resetCount"] == 1
    refreshed = client.get("/api/recovery").get_json()["data"]["items"][0]
    assert refreshed["status"] == "PENDING"
    assert refreshed["recovery_code"] == "CODE02"


def test_exception_requires_remark(client):
    client.post(
        "/api/recovery/import",
        data={"operator": "曾芸芸", "file": (recovery_workbook(client), "ADM.xlsx")},
        content_type="multipart/form-data",
    )
    item_id = client.get("/api/recovery").get_json()["data"]["items"][0]["id"]
    response = client.patch(
        f"/api/recovery/{item_id}",
        json={"status": "EXCEPTION", "remark": "", "handler": "曾芸芸"},
    )
    assert response.status_code == 400


def test_manual_code_edit_resets_to_pending_and_can_delete(client):
    client.post(
        "/api/recovery/import",
        data={"operator": "曾芸芸", "file": (recovery_workbook(client, "WRONG01"), "ADM.xlsx")},
        content_type="multipart/form-data",
    )
    item = client.get("/api/recovery").get_json()["data"]["items"][0]
    client.patch(
        f"/api/recovery/{item['id']}",
        json={"status": "COMPLETED", "remark": "原编码处理过", "handler": "曾芸芸"},
    )

    changed = client.patch(
        f"/api/recovery/{item['id']}/code",
        json={"recoveryCode": "correct02", "handler": "曾芸芸"},
    )
    assert changed.status_code == 200
    changed_item = changed.get_json()["data"]
    assert changed_item["recovery_code"] == "CORRECT02"
    assert changed_item["status"] == "PENDING"
    assert changed_item["remark"] == ""

    deleted = client.delete(
        f"/api/recovery/{item['id']}",
        json={"handler": "曾芸芸"},
    )
    assert deleted.status_code == 200
    assert client.get("/api/recovery").get_json()["data"]["pagination"]["total"] == 0
