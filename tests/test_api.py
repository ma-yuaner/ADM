from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook


def test_health(client):
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.get_json()["data"]["status"] == "UP"
    assert response.get_json()["data"]["dataMode"] == "mock"
    assert response.get_json()["data"]["financeDatabase"] == "MOCK"


def test_team_tasks_and_summary(client):
    response = client.get("/api/tasks?scope=team&person=黄娜娟&page=1&pageSize=20")
    assert response.status_code == 200
    data = response.get_json()["data"]
    assert data["pagination"]["total"] == 6
    assert data["summary"]["currentOpen"] == 6
    assert data["summary"]["unassigned"] == 2
    assert all(item["owner"] == "黄娜娟" for item in data["items"])


def test_my_tasks_respect_actual_owner(client):
    response = client.get("/api/tasks?scope=mine&person=马远尔")
    data = response.get_json()["data"]
    assert data["pagination"]["total"] == 2
    assert all(item["primaryOwner"] == "马远尔" for item in data["items"])


def test_tasks_filter_by_inclusive_creation_date_range(client, repository):
    for index, row in enumerate(repository.rows):
        row["create_time"] = datetime(2026, 9, 10 + index, 12, 0)
    repository.rows[1]["create_time"] = datetime(2026, 9, 11, 23, 59, 59)

    response = client.get(
        "/api/tasks?scope=all&createdFrom=2026-09-10&createdTo=2026-09-11"
    )
    assert response.status_code == 200
    data = response.get_json()["data"]
    assert data["pagination"]["total"] == 2
    assert {item["id"] for item in data["items"]} == {101, 102}


def test_tasks_reject_reversed_or_invalid_creation_dates(client):
    reversed_range = client.get(
        "/api/tasks?scope=all&createdFrom=2026-09-12&createdTo=2026-09-11"
    )
    assert reversed_range.status_code == 400
    assert "开始日期不能晚于结束日期" in reversed_range.get_json()["message"]

    invalid_date = client.get("/api/tasks?scope=all&createdFrom=2026-02-30")
    assert invalid_date.status_code == 400
    assert "格式无效" in invalid_date.get_json()["message"]


def test_export_respects_creation_date_range(client, repository):
    for index, row in enumerate(repository.rows):
        row["create_time"] = datetime(2026, 9, 10 + index, 12, 0)
    response = client.get(
        "/api/export?scope=all&createdFrom=2026-09-12&createdTo=2026-09-12"
    )
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), data_only=True)
    sheet = workbook["ADM待处理"]
    assert sheet.max_row == 5
    assert sheet["A5"].value == "ADM-260909-003"
    workbook.close()


def test_assign_updates_mock_repository(client):
    response = client.post("/api/assign", json={
        "admIds": [101, 104],
        "actualOwner": "李志君",
        "actor": "黄娜娟",
    })
    assert response.status_code == 200
    assert response.get_json()["data"]["updatedCount"] == 2

    tasks = client.get("/api/tasks?scope=mine&person=李志君").get_json()["data"]["items"]
    assert {101, 104}.issubset({item["id"] for item in tasks})
    transferred = next(item for item in tasks if item["id"] == 101)
    assert transferred["transferCount"] == 1
    assert transferred["lastTransferTo"] == "李志君"
    assert transferred["transferStatus"] == "转单后未接单"


def test_assign_rejects_invalid_ids(client):
    response = client.post("/api/assign", json={
        "admIds": [],
        "actualOwner": "李志君",
        "actor": "黄娜娟",
    })
    assert response.status_code == 400


def test_export_returns_valid_workbook(client):
    response = client.get("/api/export?scope=team&person=黄娜娟&source=CTRIP")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(BytesIO(response.data))
    assert workbook.sheetnames == ["ADM待处理"]
    sheet = workbook["ADM待处理"]
    assert sheet.max_row >= 5
    assert sheet["A1"].value == "ADM未结案订单核实清单"
    assert sheet.freeze_panes == "A5"
    assert sheet["A1"].fill.fgColor.rgb.endswith("17324D")
    headers = [cell.value for cell in sheet[4]]
    assert headers == [
        "ADM单号",
        "平台",
        "航司",
        "当前阶段",
        "票号",
        "ADM总金额",
        "主币种",
        "供应下发日期",
        "ADM最晚时限",
        "差异说明",
        "当前责任人",
        "是否确认",
        "实际责任人",
        "最近转单时间",
        "转单前责任人",
        "转单次数",
        "转单状态",
        "处理进度",
        "申诉状态",
        "申诉原因",
        "申诉结果",
        "结案处理结果",
        "系统-PCC：恢复编码",
    ]
    assert sheet["L5"].value is None
    assert sheet["R5"].value in {
        "已录入差异", "无差异单"
    }
    assert len(sheet.data_validations.dataValidation) == 3
    assert sheet["J5"].fill.fgColor.rgb.endswith("FFF2B2")
    assert sheet["M5"].fill.fgColor.rgb.endswith("FFF2B2")
    assert sheet["S5"].fill.fgColor.rgb.endswith("FFF2B2")
    assert sheet["T5"].value is None
    assert sheet["W5"].value is None
    assert sheet["W5"].fill.fgColor.rgb.endswith("FFF2B2")
    assert sheet["W5"].number_format == "@"


def test_wecom_send_is_disabled_by_default(client):
    response = client.post("/api/wecom/send", json={"person": "黄娜娟"})
    assert response.status_code == 503
    assert "尚未启用" in response.get_json()["message"]
