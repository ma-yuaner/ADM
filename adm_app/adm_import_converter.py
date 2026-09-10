from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .domain import appeal_submission_name
from .errors import AppError


IMPORT_COLUMNS = [
    ("id", "主键ID"),
    ("adm_no", "ADM单号（业务唯一标识）"),
    ("ota_code", "OTA平台"),
    ("ota_order_no", "OTA订单号"),
    ("lock_flag", "锁状态"),
    ("lock_operator_id", "加锁人id"),
    ("lock_operator", "加锁人"),
    ("supplier_code", "供应商"),
    ("airline", "航司（多个逗号隔开）"),
    ("supplier_type_name", "类型"),
    ("ticket_no", "票号"),
    ("ticket_count", "票号数量（冗余统计，对应adm_details行数，方便列表展示）"),
    ("total_amount", "ADM总金额（冗余汇总，= SUM(adm_details.amount)，方便列表排序展示）"),
    ("currency", "主币种（CNY/USD/EUR等，明细行可不同币种时以明细为准）"),
    ("supply_issue_date", "供应下发日期"),
    ("adm_deadline", "ADM最晚时限（回复截至时间）"),
    ("category_name", "大类"),
    ("diff_type_name", "差异类型"),
    ("diff_detail_reason", "细分差异原因（可备注，自由文本描述具体差异情况）"),
    ("owner", "当前责任人（处理人，中间可转手，历史记录见flow_logs）"),
    ("actual_owner", "实际责任人"),
    ("adm_status_name", "处理状态"),
    ("appeal_status_name", "申诉状态"),
    ("appeal_result_name", "申诉结果"),
    ("appeal_reason", "申诉原因（status=申诉时必填）"),
    ("resolution", "结案处理结果（status=已结案时必填，如：已退款¥XXX / 差异已消化 / 关单无差异等）"),
    ("handle_time", "处理时间"),
    ("auditor_name", "审核人"),
    ("auditor_time", "审核时间"),
    ("auditor_remark", "审核备注"),
    ("actual_handler", "实际处理人"),
    ("file_url", "附件地址"),
    ("status", "状态：0-无效(删除),1-有效(正常),2-暂停"),
    ("remark", "备注"),
    ("create_time", "创建时间"),
    ("update_time", "更新时间"),
    ("create_user_name", "创建人"),
    ("update_user_name", "修改人"),
]

WORKBENCH_REQUIRED_HEADERS = {"ADM单号"}
WORKBENCH_FIELDS = {
    "difference_description": "差异说明",
    "confirmation": "是否确认",
    "actual_owner": "实际责任人",
    "appeal_submission_status": "申诉状态",
    "appeal_reason": "申诉原因",
    "appeal_result": "申诉结果",
    "resolution": "结案处理结果",
    "recovery_code": "恢复编码",
}

SUPPLIER_TYPE_NAMES = {1: "平台", 2: "航司", 3: "供应商"}
CATEGORY_NAMES = {0: "票务", 1: "业务", 2: "客服"}
DIFF_TYPE_NAMES = {
    0: "订位错误",
    2: "出票错误",
    3: "税费错误",
    4: "返点错误",
    5: "退票错误",
    6: "改签错误",
    7: "取消率过高",
    8: "支付方式错误",
    9: "其他",
}
ADM_STATUS_NAMES = {0: "待申诉", 1: "申诉中", 2: "待审核", 3: "已结案"}
APPEAL_STATUS_NAMES = {0: "申诉", 1: "不申诉"}
APPEAL_RESULT_NAMES = {0: "申诉成功", 1: "申诉失败"}

FIELD_LIMITS = {
    "diff_detail_reason": 800,
    "actual_owner": 64,
    "appeal_reason": 800,
    "resolution": 800,
    "remark": 200,
}


def _excel_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _enum_name(value, names: dict[int, str]) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, str):
        stripped = value.strip()
        if stripped in names.values():
            return stripped
        try:
            value = int(stripped)
        except ValueError:
            return stripped
    return names.get(int(value), str(value))


@dataclass(frozen=True)
class WorkbenchRow:
    row_number: int
    adm_no: str
    values: dict[str, str]
    context: dict[str, str]


class AdmImportConverter:
    """把业务填写后的工作台Excel转换成ADM系统完整导入模板。"""

    def parse_workbench(self, content: bytes) -> list[WorkbenchRow]:
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise AppError("Excel文件无法读取，请上传工作台导出的.xlsx文件") from error

        try:
            sheet = workbook["ADM待处理"] if "ADM待处理" in workbook.sheetnames else workbook.active
            header_row = None
            header_map: dict[str, int] = {}
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1
            ):
                current = {_excel_text(value): index for index, value in enumerate(row) if _excel_text(value)}
                if WORKBENCH_REQUIRED_HEADERS.issubset(current):
                    header_row = row_number
                    header_map = current
                    break
            if header_row is None:
                raise AppError("Excel前10行中未找到“ADM单号”表头")

            rows: list[WorkbenchRow] = []
            seen: dict[str, WorkbenchRow] = {}
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
            ):
                adm_index = header_map["ADM单号"]
                adm_no = _excel_text(row[adm_index] if adm_index < len(row) else None)
                if not adm_no:
                    if any(_excel_text(value) for value in row):
                        raise AppError(f"第{row_number}行存在内容但ADM单号为空")
                    continue

                values = {
                    key: _excel_text(row[header_map[label]] if label in header_map and header_map[label] < len(row) else None)
                    for key, label in WORKBENCH_FIELDS.items()
                }
                context = {
                    label: _excel_text(row[index] if index < len(row) else None)
                    for label, index in header_map.items()
                    if label in {
                        "当前阶段", "是否确认", "最近转单时间", "转单前责任人",
                        "转单次数", "转单状态", "处理进度", "恢复编码",
                    }
                }
                parsed = WorkbenchRow(row_number, adm_no, values, context)
                if adm_no in seen:
                    previous = seen[adm_no]
                    if previous.values != parsed.values or previous.context != parsed.context:
                        raise AppError(
                            f"ADM {adm_no}在第{previous.row_number}行和第{row_number}行存在不同内容，请只保留一行"
                        )
                    continue
                seen[adm_no] = parsed
                rows.append(parsed)

            if not rows:
                raise AppError("Excel中没有可转换的ADM数据")
            if len(rows) > 10000:
                raise AppError("单次最多转换10000张ADM")
            return rows
        finally:
            workbook.close()

    def convert(self, rows: Iterable[WorkbenchRow], source_rows: dict[str, dict]) -> BytesIO:
        parsed_rows = list(rows)
        missing = [row.adm_no for row in parsed_rows if row.adm_no not in source_rows]
        if missing:
            preview = "、".join(missing[:20])
            suffix = f"等{len(missing)}张" if len(missing) > 20 else ""
            raise AppError(f"数据库中未找到有效ADM：{preview}{suffix}，本次未生成导入文件")

        output_rows = [self._merge_row(row, source_rows[row.adm_no]) for row in parsed_rows]
        self._validate_rows(output_rows)
        return self._build_workbook(output_rows)

    @staticmethod
    def _submission_status_from_source(source: dict) -> str:
        return appeal_submission_name(source)

    def _merge_row(self, workbench: WorkbenchRow, source: dict) -> dict:
        merged = dict(source)
        values = workbench.values

        if values["difference_description"]:
            merged["diff_detail_reason"] = values["difference_description"]
        if values["actual_owner"]:
            merged["actual_owner"] = values["actual_owner"]
        if values["appeal_reason"]:
            merged["appeal_reason"] = values["appeal_reason"]
        if values["resolution"]:
            merged["resolution"] = values["resolution"]

        submission = values["appeal_submission_status"]
        source_submission = self._submission_status_from_source(source)
        if submission and submission != source_submission:
            if submission == "待提交":
                merged["adm_status"] = 0
                merged["appeal_status"] = None
                merged["appeal_result"] = None
            elif submission == "已提交":
                merged["adm_status"] = 1
                merged["appeal_status"] = 0
                merged["appeal_result"] = None
            elif submission == "不提交可结案":
                merged["adm_status"] = 2
                merged["appeal_status"] = 1
                merged["appeal_result"] = None
            else:
                raise AppError(
                    f"第{workbench.row_number}行ADM {workbench.adm_no}的申诉状态“{submission}”无效"
                )

        appeal_result = values["appeal_result"]
        if appeal_result:
            if appeal_result not in {"申诉成功", "申诉失败"}:
                raise AppError(
                    f"第{workbench.row_number}行ADM {workbench.adm_no}的申诉结果“{appeal_result}”无效"
                )
            merged["adm_status"] = 2
            merged["appeal_status"] = 0
            merged["appeal_result"] = 0 if appeal_result == "申诉成功" else 1

        merged["remark"] = self._merge_remark(_excel_text(source.get("remark")), workbench.context)
        self._add_display_values(merged)
        return merged

    @staticmethod
    def _merge_remark(original: str, context: dict[str, str]) -> str:
        compact = []
        aliases = {
            "当前阶段": "阶段",
            "是否确认": "确认",
            "最近转单时间": "转单时间",
            "转单前责任人": "转出人",
            "转单次数": "转单数",
            "转单状态": "转单状态",
            "处理进度": "进度",
            "恢复编码": "恢复编码",
        }
        for label, alias in aliases.items():
            value = context.get(label, "")
            if value:
                compact.append(f"{alias}={value}")
        generated = f"[工作台回传]{'；'.join(compact)}" if compact else ""

        base = original
        marker = "[工作台回传]"
        if marker in base:
            base = base.split(marker, 1)[0].rstrip("； ")
        return "；".join(value for value in (base, generated) if value)

    @staticmethod
    def _add_display_values(row: dict) -> None:
        row["supplier_type_name"] = _enum_name(row.get("supplier_type"), SUPPLIER_TYPE_NAMES)
        row["category_name"] = _enum_name(row.get("category"), CATEGORY_NAMES)
        row["diff_type_name"] = _enum_name(row.get("diff_type"), DIFF_TYPE_NAMES)
        row["adm_status_name"] = _enum_name(row.get("adm_status"), ADM_STATUS_NAMES)
        row["appeal_status_name"] = _enum_name(row.get("appeal_status"), APPEAL_STATUS_NAMES)
        row["appeal_result_name"] = _enum_name(row.get("appeal_result"), APPEAL_RESULT_NAMES)

    @staticmethod
    def _validate_rows(rows: list[dict]) -> None:
        errors = []
        for row in rows:
            adm_no = _excel_text(row.get("adm_no"))
            for key, limit in FIELD_LIMITS.items():
                value = _excel_text(row.get(key))
                if len(value) > limit:
                    errors.append(f"ADM {adm_no}的{key}为{len(value)}字，超过{limit}字")
            if len(errors) >= 20:
                break
        if errors:
            raise AppError("；".join(errors) + "。请精简后重新转换")

    @staticmethod
    def _cell_value(value):
        if isinstance(value, (date, datetime)):
            return value
        return "" if value is None else value

    def _build_workbook(self, rows: list[dict]) -> BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = f"ADM跟单-{datetime.now():%Y%m%d-%H%M}"
        headers = [label for _, label in IMPORT_COLUMNS]
        sheet.append(headers)

        border = Border(
            left=Side(style="thin", color="B7B7B7"),
            right=Side(style="thin", color="B7B7B7"),
            top=Side(style="thin", color="B7B7B7"),
            bottom=Side(style="thin", color="B7B7B7"),
        )
        for cell in sheet[1]:
            cell.font = Font(name="宋体", size=14, bold=True, color="000000")
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        sheet.row_dimensions[1].height = 48

        for source in rows:
            sheet.append([self._cell_value(source.get(key)) for key, _ in IMPORT_COLUMNS])

        date_columns = {15}
        datetime_columns = {16, 27, 29, 35, 36}
        numeric_columns = {1, 5, 6, 12, 13, 33}
        for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            for column_number, cell in enumerate(row, start=1):
                cell.font = Font(name="宋体", size=10, color="000000")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border
                if column_number in date_columns:
                    cell.number_format = "yyyy-mm-dd"
                elif column_number in datetime_columns:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif column_number == 13:
                    cell.number_format = "#,##0.00"
                elif column_number in numeric_columns:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            sheet.row_dimensions[row_number].height = 26

        widths = {
            1: 12, 2: 24, 3: 14, 4: 20, 5: 10, 6: 12, 7: 14, 8: 16,
            9: 17, 10: 11, 11: 22, 12: 17, 13: 18, 14: 12, 15: 14, 16: 20,
            17: 11, 18: 16, 19: 34, 20: 18, 21: 16, 22: 13, 23: 13,
            24: 13, 25: 34, 26: 34, 27: 20, 28: 13, 29: 20, 30: 24,
            31: 15, 32: 28, 33: 15, 34: 42, 35: 20, 36: 20, 37: 14, 38: 14,
        }
        for column_number, width in widths.items():
            sheet.column_dimensions[get_column_letter(column_number)].width = width

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:AL{max(1, sheet.max_row)}"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:1"

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream
