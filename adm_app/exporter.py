from __future__ import annotations

import json
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_COLUMNS = [
    {"key": "admNo", "label": "ADM单号"},
    {"key": "otaCode", "label": "平台"},
    {"key": "airline", "label": "航司"},
    {"key": "stageName", "label": "当前阶段"},
    {"key": "ticketNo", "label": "票号"},
    {"key": "amount", "label": "ADM总金额"},
    {"key": "currency", "label": "主币种"},
    {"key": "supplyIssueDate", "label": "供应下发日期"},
    {"key": "deadline", "label": "ADM最晚时限"},
    {"key": "differenceDescription", "label": "差异说明"},
    {"key": "owner", "label": "当前责任人"},
    {"key": "confirmation", "label": "是否确认"},
    {"key": "actualOwner", "label": "实际责任人"},
    {"key": "lastTransferTime", "label": "最近转单时间"},
    {"key": "lastTransferFrom", "label": "转单前责任人"},
    {"key": "transferCount", "label": "转单次数"},
    {"key": "transferStatus", "label": "转单状态"},
    {"key": "handlingProgress", "label": "处理进度"},
    {"key": "appealSubmissionStatus", "label": "申诉状态"},
    {"key": "appealReason", "label": "申诉原因"},
    {"key": "appealResultName", "label": "申诉结果"},
    {"key": "resolution", "label": "结案处理结果"},
    {"key": "recoveryCode", "label": "恢复编码"},
]


class ExcelExportService:
    def __init__(self, profiles_file: Path):
        self.profiles_file = profiles_file
        self.profiles = self._load_profiles()

    def _load_profiles(self) -> dict:
        if not self.profiles_file.exists():
            return {"DEFAULT": {"columns": DEFAULT_COLUMNS}}
        with self.profiles_file.open("r", encoding="utf-8") as handle:
            configured = json.load(handle)
        configured.setdefault("DEFAULT", {"columns": DEFAULT_COLUMNS})
        return configured

    def _columns(self) -> list[dict]:
        return self.profiles["DEFAULT"].get("columns", DEFAULT_COLUMNS)

    @staticmethod
    def _cell_value(value):
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return value
        return "" if value is None else value

    def build(
        self,
        tasks: list[dict],
        title: str = "ADM未结案订单核实清单",
        subtitle: str | None = None,
    ) -> BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ADM待处理"
        columns = self._columns()
        last_column = get_column_letter(len(columns))
        subtitle = subtitle or f"共{len(tasks)}张｜导出时间：{datetime.now():%Y-%m-%d %H:%M}"

        sheet.merge_cells(f"A1:{last_column}1")
        sheet["A1"] = title
        sheet["A1"].font = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="17324D")
        sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[1].height = 34

        sheet.merge_cells(f"A2:{last_column}2")
        sheet["A2"] = subtitle
        sheet["A2"].font = Font(name="微软雅黑", size=10, color="53657A")
        sheet["A2"].fill = PatternFill("solid", fgColor="EAF0F6")
        sheet["A2"].alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[2].height = 24

        header_row = 4
        data_start_row = 5
        sheet.append([])
        sheet.append([column["label"] for column in columns])
        thin_border = Border(
            left=Side(style="thin", color="D7E0EA"),
            right=Side(style="thin", color="D7E0EA"),
            top=Side(style="thin", color="D7E0EA"),
            bottom=Side(style="thin", color="D7E0EA"),
        )
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2457D6")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        sheet.row_dimensions[header_row].height = 32

        for item in tasks:
            sheet.append([self._cell_value(item.get(column["key"])) for column in columns])

        editable_keys = {
            "differenceDescription", "confirmation", "actualOwner", "handlingProgress",
            "appealSubmissionStatus", "appealReason", "appealResultName", "resolution",
            "recoveryCode",
        }
        editable_columns = {
            column["key"]: index
            for index, column in enumerate(columns, start=1)
            if column["key"] in editable_keys
        }
        input_column = editable_columns["confirmation"]
        transfer_status_column = next(
            (
                index for index, column in enumerate(columns, start=1)
                if column["key"] == "transferStatus"
            ),
            None,
        )
        recovery_code_column = next(
            (
                index for index, column in enumerate(columns, start=1)
                if column["key"] == "recoveryCode"
            ),
            None,
        )
        for row_index, (item, row) in enumerate(
            zip(tasks, sheet.iter_rows(min_row=data_start_row)),
            start=data_start_row,
        ):
            alert_level = item.get("alertLevel")
            if alert_level == "P0":
                row_fill = "FDECEC"
            elif alert_level == "P1":
                row_fill = "FFF4D6"
            else:
                row_fill = "F7FAFC" if row_index % 2 == 0 else "FFFFFF"
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(name="微软雅黑", size=10, color="26384A")
                cell.fill = PatternFill("solid", fgColor=row_fill)
                cell.border = thin_border
            for editable_column in editable_columns.values():
                row[editable_column - 1].fill = PatternFill("solid", fgColor="FFF2B2")
            row[input_column - 1].alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            if recovery_code_column:
                recovery_cell = row[recovery_code_column - 1]
                recovery_cell.fill = PatternFill("solid", fgColor="FFF2B2")
                recovery_cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=False
                )
            if transfer_status_column and item.get("transferStatus") == "转单后未接单":
                transfer_cell = row[transfer_status_column - 1]
                transfer_cell.fill = PatternFill("solid", fgColor="FFF2B2")
                transfer_cell.font = Font(name="微软雅黑", size=10, bold=True, color="9D6200")
            sheet.row_dimensions[row_index].height = 30

        sheet.freeze_panes = f"A{data_start_row}"
        sheet.auto_filter.ref = f"A{header_row}:{last_column}{max(header_row, sheet.max_row)}"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = f"1:{header_row}"
        confirmation_validation = DataValidation(
            type="list",
            formula1='"确认,非我司订单,资料不全"',
            allow_blank=True,
        )
        confirmation_validation.promptTitle = "请选择核实结果"
        confirmation_validation.prompt = "确认、非我司订单或资料不全"
        confirmation_validation.error = "请从下拉选项中选择"
        confirmation_validation.errorTitle = "填写内容不正确"
        confirmation_validation.showErrorMessage = True
        sheet.add_data_validation(confirmation_validation)
        confirmation_validation.add(
            f"{get_column_letter(input_column)}{data_start_row}:"
            f"{get_column_letter(input_column)}{max(data_start_row, sheet.max_row)}"
        )
        appeal_status_column = editable_columns.get("appealSubmissionStatus")
        if appeal_status_column:
            appeal_status_validation = DataValidation(
                type="list",
                formula1='"待提交,已提交,不提交可结案"',
                allow_blank=True,
            )
            appeal_status_validation.promptTitle = "请选择申诉状态"
            appeal_status_validation.prompt = "待提交、已提交或不提交可结案"
            appeal_status_validation.error = "请从下拉选项中选择"
            appeal_status_validation.showErrorMessage = True
            sheet.add_data_validation(appeal_status_validation)
            appeal_status_validation.add(
                f"{get_column_letter(appeal_status_column)}{data_start_row}:"
                f"{get_column_letter(appeal_status_column)}{max(data_start_row, sheet.max_row)}"
            )
        appeal_result_column = editable_columns.get("appealResultName")
        if appeal_result_column:
            appeal_result_validation = DataValidation(
                type="list",
                formula1='"申诉成功,申诉失败"',
                allow_blank=True,
            )
            appeal_result_validation.promptTitle = "请选择申诉结果"
            appeal_result_validation.prompt = "申诉成功或申诉失败"
            appeal_result_validation.error = "请从下拉选项中选择"
            appeal_result_validation.showErrorMessage = True
            sheet.add_data_validation(appeal_result_validation)
            appeal_result_validation.add(
                f"{get_column_letter(appeal_result_column)}{data_start_row}:"
                f"{get_column_letter(appeal_result_column)}{max(data_start_row, sheet.max_row)}"
            )
        preferred_widths = {
            "admNo": 20, "otaCode": 13, "airline": 10, "stageName": 17,
            "ticketNo": 20, "amount": 14, "currency": 10,
            "supplyIssueDate": 14, "deadline": 19, "differenceDescription": 32,
            "owner": 14, "confirmation": 15, "actualOwner": 14,
            "lastTransferTime": 19, "lastTransferFrom": 14,
            "transferCount": 10, "transferStatus": 17,
            "handlingProgress": 15, "appealSubmissionStatus": 18,
            "appealReason": 32, "appealResultName": 14, "resolution": 32,
            "recoveryCode": 18,
        }
        for column_index, column in enumerate(columns, start=1):
            width = preferred_widths.get(column["key"], 15)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
            key = column["key"]
            if key in {"deadline", "lastTransferTime", "updateTime"}:
                for cell in sheet.iter_cols(min_col=column_index, max_col=column_index, min_row=data_start_row):
                    cell[0].number_format = "yyyy-mm-dd hh:mm"
            elif key == "supplyIssueDate":
                for cell in sheet.iter_cols(min_col=column_index, max_col=column_index, min_row=data_start_row):
                    cell[0].number_format = "yyyy-mm-dd"
            elif key == "amount":
                for cell in sheet.iter_cols(min_col=column_index, max_col=column_index, min_row=data_start_row):
                    cell[0].number_format = "#,##0.00"
                    cell[0].alignment = Alignment(horizontal="right", vertical="top")
            elif key == "recoveryCode":
                for cell in sheet.iter_cols(min_col=column_index, max_col=column_index, min_row=data_start_row):
                    cell[0].number_format = "@"

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream
