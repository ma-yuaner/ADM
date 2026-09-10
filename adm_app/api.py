from __future__ import annotations

from datetime import datetime
from io import BytesIO

from flask import Blueprint, current_app, jsonify, request, send_file
from openpyxl import load_workbook

from .adm_import_converter import AdmImportConverter
from .errors import AppError
from .exporter import ExcelExportService
from .json_utils import json_ready


api = Blueprint("api", __name__, url_prefix="/api")


def _repository():
    return current_app.extensions["adm_repository"]


def _wecom():
    return current_app.extensions["wecom_service"]


def _recovery():
    return current_app.extensions["recovery_store"]


def _success(data=None, message="success", status=200):
    return jsonify(json_ready({"success": True, "message": message, "data": data})), status


def _clean_text(name: str, max_length: int, required: bool = False) -> str:
    value = (request.args.get(name) or "").strip()
    if required and not value:
        raise AppError(f"{name}不能为空")
    if len(value) > max_length:
        raise AppError(f"{name}长度不能超过{max_length}")
    return value


def _positive_int(name: str, default: int, maximum: int) -> int:
    raw = request.args.get(name, str(default))
    try:
        value = int(raw)
    except ValueError as error:
        raise AppError(f"{name}必须是整数") from error
    if value < 1 or value > maximum:
        raise AppError(f"{name}必须在1到{maximum}之间")
    return value


def _excel_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _task_filters(export: bool = False) -> dict:
    return {
        "scope": _clean_text("scope", 10) or "mine",
        "person": _clean_text("person", 64),
        "source": _clean_text("source", 64),
        "alert": _clean_text("alert", 16),
        "stage": _clean_text("stage", 50),
        "search": _clean_text("search", 100),
        "page": 1 if export else _positive_int("page", 1, 100000),
        "page_size": 10000 if export else _positive_int("pageSize", 20, 200),
    }


@api.get("/health")
def health():
    state = _repository().health()
    return _success({
        "status": "UP",
        "dataMode": _repository().mode,
        "database": state["database"],
        "financeDatabase": state.get("financeDatabase", "DISABLED"),
        "writeEnabled": state["writeEnabled"],
        "time": datetime.now(),
    })


@api.get("/config")
def config():
    return _success({
        "defaultPerson": current_app.config["DEFAULT_PERSON"],
        "dataMode": _repository().mode,
        "writeEnabled": _repository().health()["writeEnabled"],
        "wecomEnabled": _wecom().enabled,
    })


@api.get("/people")
def people():
    return _success({"items": _repository().people()})


@api.get("/tasks")
def tasks():
    return _success(_repository().list_tasks(_task_filters()))


@api.post("/assign")
def assign():
    body = request.get_json(silent=True) or {}
    ids = body.get("admIds")
    if not isinstance(ids, list) or not ids or len(ids) > 500:
        raise AppError("admIds必须是包含1到500个ID的数组")
    try:
        adm_ids = list(dict.fromkeys(int(value) for value in ids))
    except (TypeError, ValueError) as error:
        raise AppError("admIds只能包含整数") from error
    if any(value <= 0 for value in adm_ids):
        raise AppError("admIds只能包含正整数")

    actual_owner = str(body.get("actualOwner") or "").strip()
    actor = str(body.get("actor") or "").strip()
    if not actual_owner or len(actual_owner) > 64:
        raise AppError("actualOwner不能为空且长度不能超过64")
    if not actor or len(actor) > 64:
        raise AppError("actor不能为空且长度不能超过64")
    client_ip = (request.access_route[0] if request.access_route else request.remote_addr) or "0.0.0.0"
    result = _repository().assign(adm_ids, actual_owner, actor, client_ip)
    return _success(result, f"成功分配{result['updatedCount']}张ADM")


@api.get("/export")
def export_tasks():
    result = _repository().list_tasks(_task_filters(export=True))
    if result["pagination"]["total"] > 10000:
        raise AppError("单次导出不能超过10000张ADM，请增加筛选条件")
    exporter = ExcelExportService(current_app.config["EXPORT_PROFILES_FILE"])
    stream = exporter.build(result["items"])
    person = _clean_text("person", 64) or "全部人员"
    source = _clean_text("source", 64) or "全部数据源"
    filename = f"ADM待处理_{source}_{person}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    return send_file(
        stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        max_age=0,
    )


@api.post("/wecom/send")
def send_wecom():
    body = request.get_json(silent=True) or {}
    person = str(body.get("person") or "").strip()
    if not person or len(person) > 64:
        raise AppError("person不能为空且长度不能超过64")
    result = _repository().list_tasks({
        "scope": "mine",
        "person": person,
        "source": "",
        "alert": "",
        "stage": "",
        "search": "",
        "page": 1,
        "page_size": 10000,
    })
    tasks = result["items"]
    if not tasks:
        raise AppError(f"{person}当前没有未结案ADM")
    exporter = ExcelExportService(current_app.config["EXPORT_PROFILES_FILE"])
    workbook = exporter.build(
        tasks,
        title=f"{person}—ADM未结案订单核实清单",
        subtitle=f"共{len(tasks)}张｜请核实订单归属、差异及申诉进度",
    )
    filename = f"ADM未结案核实_{person}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    return _success(
        _wecom().send_excel(person, workbook, filename, len(tasks)),
        f"已发送{len(tasks)}张未结案ADM给{person}",
    )


@api.post("/recovery/import")
def import_recovery_codes():
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        raise AppError("请选择需要导入的Excel文件")
    if not uploaded.filename.lower().endswith(".xlsx"):
        raise AppError("只支持.xlsx文件")
    imported_by = (request.form.get("operator") or "").strip()
    if not imported_by or len(imported_by) > 64:
        raise AppError("导入操作人不能为空且长度不能超过64")

    try:
        workbook = load_workbook(BytesIO(uploaded.read()), read_only=True, data_only=True)
    except Exception as error:
        raise AppError("Excel文件无法读取，请使用工作台导出的原始模板") from error
    try:
        sheet = workbook["ADM待处理"] if "ADM待处理" in workbook.sheetnames else workbook.active
        header_row = None
        header_map: dict[str, int] = {}
        recovery_header = ""
        for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            current = {_excel_text(value): index for index, value in enumerate(row)}
            recovery_header = "恢复编码" if "恢复编码" in current else ("编码" if "编码" in current else "")
            if "ADM单号" in current and recovery_header:
                header_row = row_number
                header_map = current
                break
        if header_row is None:
            raise AppError("Excel缺少“ADM单号”或“恢复编码”列")

        pending: dict[str, str] = {}
        invalid_rows: list[dict] = []
        blank_count = 0
        duplicate_count = 0
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
        ):
            adm_no = _excel_text(row[header_map["ADM单号"]] if header_map["ADM单号"] < len(row) else None)
            recovery_code = _excel_text(
                row[header_map[recovery_header]] if header_map[recovery_header] < len(row) else None
            ).upper()
            if not adm_no and not recovery_code:
                continue
            if not recovery_code:
                blank_count += 1
                continue
            if not adm_no:
                invalid_rows.append({"row": row_number, "reason": "恢复编码有值但ADM单号为空"})
                continue
            if len(recovery_code) > 64:
                invalid_rows.append({"row": row_number, "admNo": adm_no, "reason": "恢复编码超过64位"})
                continue
            if adm_no in pending:
                if pending[adm_no] != recovery_code:
                    invalid_rows.append({"row": row_number, "admNo": adm_no, "reason": "同一ADM存在不同恢复编码"})
                else:
                    duplicate_count += 1
                continue
            pending[adm_no] = recovery_code
        if len(pending) > 10000:
            raise AppError("单次最多导入10000张ADM")

        tasks = _repository().find_by_adm_numbers(list(pending))
        missing = sorted(set(pending) - set(tasks))
        counters = {"created": 0, "reset": 0, "unchanged": 0}
        for adm_no, task in tasks.items():
            result = _recovery().upsert(task, pending[adm_no], imported_by)
            counters[result] += 1
        return _success({
            "validCodeCount": len(pending),
            "importedCount": len(tasks),
            "createdCount": counters["created"],
            "resetCount": counters["reset"],
            "unchangedCount": counters["unchanged"],
            "blankCount": blank_count,
            "duplicateCount": duplicate_count,
            "missingAdmNumbers": missing[:100],
            "invalidRows": invalid_rows[:100],
        }, f"成功导入{len(tasks)}条恢复编码")
    finally:
        workbook.close()


@api.post("/adm-import/convert")
def convert_adm_import_workbook():
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        raise AppError("请选择业务填写后的工作台Excel")
    if not uploaded.filename.lower().endswith(".xlsx"):
        raise AppError("只支持.xlsx文件")

    content = uploaded.read()
    if not content:
        raise AppError("上传的Excel文件为空")
    converter = AdmImportConverter()
    workbench_rows = converter.parse_workbench(content)
    source_rows = _repository().find_import_rows_by_adm_numbers(
        [row.adm_no for row in workbench_rows]
    )
    recovery_rows = _recovery().find_by_adm_numbers(
        [row.adm_no for row in workbench_rows]
    )
    stream = converter.convert(workbench_rows, source_rows, recovery_rows)
    filename = f"ADM管理导入_{datetime.now():%Y%m%d_%H%M}.xlsx"
    return send_file(
        stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        max_age=0,
    )


@api.get("/recovery")
def recovery_items():
    status = _clean_text("status", 20)
    search = _clean_text("search", 100)
    page = _positive_int("page", 1, 100000)
    page_size = _positive_int("pageSize", 20, 200)
    return _success(_recovery().list(status, search, page, page_size))


@api.patch("/recovery/<int:item_id>")
def update_recovery(item_id: int):
    body = request.get_json(silent=True) or {}
    status = str(body.get("status") or "").strip().upper()
    remark = str(body.get("remark") or "").strip()
    handler = str(body.get("handler") or "").strip()
    if len(remark) > 500:
        raise AppError("备注不能超过500字")
    if status == "EXCEPTION" and not remark:
        raise AppError("标记异常时必须填写处理备注")
    if not handler or len(handler) > 64:
        raise AppError("处理人不能为空且长度不能超过64")
    return _success(_recovery().update(item_id, status, remark, handler), "恢复状态已更新")


@api.patch("/recovery/<int:item_id>/code")
def update_recovery_code(item_id: int):
    body = request.get_json(silent=True) or {}
    recovery_code = str(body.get("recoveryCode") or "").strip().upper()
    handler = str(body.get("handler") or "").strip()
    if not recovery_code:
        raise AppError("恢复编码不能为空；不需要该记录时请使用删除")
    if len(recovery_code) > 64:
        raise AppError("恢复编码不能超过64位")
    if not handler or len(handler) > 64:
        raise AppError("处理人不能为空且长度不能超过64")
    item = _recovery().update_code(item_id, recovery_code, handler)
    message = "恢复编码已修改并重置为待恢复" if item["codeChanged"] else "恢复编码未发生变化"
    return _success(item, message)


@api.delete("/recovery/<int:item_id>")
def delete_recovery(item_id: int):
    body = request.get_json(silent=True) or {}
    handler = str(body.get("handler") or "").strip()
    if not handler or len(handler) > 64:
        raise AppError("处理人不能为空且长度不能超过64")
    item = _recovery().delete(item_id)
    return _success({"id": item_id, "admNo": item["adm_no"], "deletedBy": handler}, "恢复编码记录已删除")
